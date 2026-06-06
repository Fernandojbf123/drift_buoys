import win32com.client
import os
from pptx import Presentation
from configs.manager_doc_config import get_ruta_a_carpeta_de_guardado_del_documento


def reemplazar_en_slide(slide, diccionario):
    for shape in slide.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    for k, v in diccionario.items():
                        if k in run.text:
                            run.text = run.text.replace(k, str(v))
        if shape.has_table:
            table = shape.table
            for row in table.rows:
                for cell in row.cells:
                    for k, v in diccionario.items():
                        if k in cell.text:
                            cell.text = cell.text.replace(k, str(v))
                            
def duplicar_slide(ppt, slide):
    layout = slide.slide_layout
    new_slide = ppt.slides.add_slide(layout)

    for shape in slide.shapes:
        new_slide.shapes._spTree.append(shape.element)
    return new_slide
                           

def generar_ppt_por_sondas(ppt, diccionario_sondas):
    template_slide = ppt.slides[0]

    slides_generados = {}

    for serial, data in diccionario_sondas.items():

        slide = duplicar_slide(ppt, template_slide)

        reemplazar_en_slide(slide, data)

        slides_generados[serial] = slide

    xml_slides = ppt.slides._sldIdLst
    ppt.slides._sldIdLst.remove(xml_slides[0])
    return ppt, slides_generados


def exportar_ppt(ppt_path, output_folder):
    ppt_app = win32com.client.Dispatch("PowerPoint.Application")
    ppt_app.Visible = 1

    presentation = ppt_app.Presentations.Open(ppt_path)

    os.makedirs(output_folder, exist_ok=True)

    pdf_path = os.path.join(output_folder, "reporte_sondas.pdf")
    presentation.SaveAs(pdf_path, 32)
    
    for i in range(1, presentation.Slides.Count + 1):
        slide = presentation.Slides(i)

        for j in range(1, slide.Shapes.Count + 1):
            pass  # placeholder si luego quieres exportar shapes individuales

        img_path = os.path.join(output_folder, f"esquema_sonda_{i}.png")
        slide.Export(img_path, "PNG")

    presentation.Close()
    ppt_app.Quit()
    
from pptx import Presentation


def ppt_sondas(
    ppt,
    diccionario_sondas,
    output_ppt="output.pptx",
    output_folder="salidas"):

    ppt, slides = generar_ppt_por_sondas(ppt, diccionario_sondas)

    ruta = get_ruta_a_carpeta_de_guardado_del_documento()
    output_ppt = os.path.join(ruta, output_ppt) 
    ppt.save(output_ppt)

    exportar_ppt(output_ppt, output_folder)
    return output_ppt, output_folder